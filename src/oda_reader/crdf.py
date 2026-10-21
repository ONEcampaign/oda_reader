"""Read the OECD Climate-Related Development Finance (CRDF) workbooks.

The OECD publishes the CRDF as Excel workbooks in a hand-maintained folder at
https://webfs.oecd.org/climate/, one subfolder per perspective. File names are
resolved from the live folder listing, and every worksheet's header row must
match one of the layouts in ``schemas/mappings/crdf_schema.json``. Anything
else raises ``CRDFSchemaError``.
"""

import difflib
import functools
import hashlib
import json
import re
import shutil
from datetime import datetime
from html import unescape
from pathlib import Path
from typing import Literal, cast
from urllib.parse import urljoin

import pandas as pd
import requests
from openpyxl import load_workbook

from oda_reader import _http_primitives
from oda_reader._cache.manager import CacheEntry, bulk_cache_manager
from oda_reader._http_primitives import (
    API_RATE_LIMITER,
    DEFAULT_HEADERS,
    _get_http_session,
    _get_same_origin,
)
from oda_reader.common import ImporterPaths, logger
from oda_reader.download.download_tools import (
    _fetch_revalidation_token,
    _finalize_resolved_url,
    _stream_to_file,
)
from oda_reader.exceptions import BulkDownloadHTTPError, CRDFSchemaError

CRDF_URL = "https://webfs.oecd.org/climate/"

# The provider folder was called DonorPerspective before the June 2026 move.
_PERSPECTIVE_FOLDERS = {
    "ProviderPerspective": "provider",
    "DonorPerspective": "provider",
    "RecipientPerspective": "recipient",
}

# One entry of the IIS directory listing: date, time, size or <dir>, link.
_LISTING_ENTRY_RE = re.compile(
    r"(\d{1,2}/\d{1,2}/\d{4}\s+\d{1,2}:\d{2} [AP]M)\s+(&lt;dir&gt;|\d+)\s+"
    r'<A HREF="([^"]+)">([^<]+)</A>',
    re.IGNORECASE,
)
# The years a workbook holds, from its name: "CRDF-RP_2014-2015.xlsx".
_FILE_YEARS_RE = re.compile(r"(\d{4})(?:-(\d{4}))?\.xlsx$")
# Constant-price amount headers embed their price year: "- 2024 USD thousand".
_PRICE_YEAR_RE = re.compile(r" - (\d{4}) USD ")
_PRICE_YEAR_PLACEHOLDER = " - {price_year} USD "

# Connect and read timeouts for a folder listing request, in seconds.
_LISTING_TIMEOUT = (10, 30)

Perspective = Literal["provider", "recipient"]


@functools.cache
def _schema() -> dict:
    with (ImporterPaths.mappings / "crdf_schema.json").open(encoding="utf-8") as f:
        return json.load(f)


def _get_listing(url: str) -> tuple[int, str]:
    """GET a folder listing past the HTTP cache, with a timeout.

    Goes through the same session, rate limiter, headers and same-origin
    redirect check as ``get_response_text``, which sets no timeout.
    """
    API_RATE_LIMITER.wait()
    if _http_primitives._CACHE_ENABLED:
        response = _get_same_origin(
            _get_http_session().get,
            url,
            DEFAULT_HEADERS,
            force_refresh=True,
            timeout=_LISTING_TIMEOUT,
        )
    else:
        response = _get_same_origin(
            requests.get, url, DEFAULT_HEADERS, timeout=_LISTING_TIMEOUT
        )
    return cast(int, response.status_code), response.text


def _list_folder(url: str) -> list[dict]:
    status, text = _get_listing(url)
    if status != 200:
        raise BulkDownloadHTTPError(status_code=status, url=url, body=text)
    entries = [
        {
            "name": unescape(name),
            "is_dir": size.lower() == "&lt;dir&gt;",
            "size_bytes": None if size.lower() == "&lt;dir&gt;" else int(size),
            "modified": datetime.strptime(" ".join(stamp.split()), "%m/%d/%Y %I:%M %p"),
            "url": _finalize_resolved_url(urljoin(url, href)),
        }
        for stamp, size, href, name in _LISTING_ENTRY_RE.findall(text)
    ]
    if not entries:
        raise CRDFSchemaError(
            source=url, detail="The folder listing matched no entries."
        )
    return entries


def list_crdf_files() -> pd.DataFrame:
    """List the workbooks in the live CRDF folder.

    Reads the folder listing at ``CRDF_URL`` and each perspective subfolder,
    always from the network.

    Returns:
        pd.DataFrame: One row per workbook with ``name``, ``perspective``
            (``"provider"``, ``"recipient"`` or missing for top-level files
            such as the imputed multilateral shares), ``start_year`` and
            ``end_year`` (parsed from the name, missing when it has no years),
            ``size_bytes``, ``modified`` (as the listing shows it, in the
            server's local time) and ``url``.
    """
    rows = []
    for entry in _list_folder(CRDF_URL):
        if not entry["is_dir"]:
            rows.append({**entry, "perspective": None})
            continue
        perspective = _PERSPECTIVE_FOLDERS.get(entry["name"])
        if perspective is None:
            logger.warning(f"Skipping unknown CRDF folder {entry['url']}")
            continue
        rows.extend(
            {**file, "perspective": perspective}
            for file in _list_folder(entry["url"])
            if not file["is_dir"]
        )

    listing = pd.DataFrame(
        rows,
        columns=[
            "name",
            "perspective",
            "start_year",
            "end_year",
            "size_bytes",
            "modified",
            "url",
        ],
    )
    years = listing["name"].str.extract(_FILE_YEARS_RE).apply(pd.to_numeric)
    listing["start_year"] = years[0].astype("Int32")
    listing["end_year"] = years[1].fillna(years[0]).astype("Int32")
    return listing


def _fetch(url: str) -> tuple[Path, str | None]:
    """Return a cached copy of the workbook at *url* and its ETag."""
    etag = _fetch_revalidation_token(url)
    key = hashlib.sha1(url.encode()).hexdigest()
    entry = CacheEntry(
        key=key,
        filename=f"{key}.xlsx",
        fetcher=lambda path: _stream_to_file(url, {}, path),
        version=etag,
    )
    return bulk_cache_manager().ensure(entry), etag


def download_crdf_file(name: str, save_to_path: Path | str) -> Path:
    """Download one workbook from the live CRDF folder.

    Args:
        name: The file name as the listing shows it, for example
            ``"CRDF-RP_2024.xlsx"`` or ``"Imputed_multilateral_shares.xlsx"``.
        save_to_path: A directory to save the file into under *name*, or the
            full path to save it to.

    Returns:
        Path: Where the file was saved.

    Raises:
        ValueError: If no file in the listing has this name.
    """
    listing = list_crdf_files()
    match = listing.loc[listing["name"] == name]
    if match.empty:
        close = difflib.get_close_matches(name, listing["name"].tolist())
        raise ValueError(
            f"{name!r} is not in the CRDF listing at {CRDF_URL}. Close names: {close}"
        )

    cached, _ = _fetch(match["url"].iloc[0])
    target = Path(save_to_path)
    if target.is_dir():
        target = target / name
    shutil.copyfile(cached, target)
    return target


def _match_layout(
    headers: list[str], *, source: str, sheet: str
) -> tuple[str, list[str], int]:
    """Find the layout whose headers equal *headers*, with the price year abstracted.

    Returns:
        The layout name, the canonical column name for each header, and the
        price year of the constant-price amounts.
    """
    if len(set(headers)) != len(headers):
        duplicated = sorted({h for h in headers if headers.count(h) > 1})
        raise CRDFSchemaError(
            source=source, sheet=sheet, detail=f"Duplicated headers {duplicated}."
        )

    normalised = [_PRICE_YEAR_RE.sub(_PRICE_YEAR_PLACEHOLDER, h) for h in headers]
    layouts = _schema()["layouts"]
    found = set(normalised)
    name = next((n for n, lay in layouts.items() if set(lay["columns"]) == found), None)
    if name is None:
        nearest = min(layouts, key=lambda n: len(set(layouts[n]["columns"]) ^ found))
        expected = layouts[nearest]["columns"]
        raise CRDFSchemaError(
            source=source,
            sheet=sheet,
            detail="The header row matches no known CRDF layout.",
            layout=nearest,
            unmapped=tuple(
                h for h, n in zip(headers, normalised, strict=True) if n not in expected
            ),
            missing=tuple(h for h in expected if h not in found),
        )

    years = {y for h in headers for y in _PRICE_YEAR_RE.findall(h)}
    if len(years) != 1:
        raise CRDFSchemaError(
            source=source,
            sheet=sheet,
            layout=name,
            detail=f"Expected one price year in the headers, found {sorted(years)}.",
        )
    columns = layouts[name]["columns"]
    return name, [columns[n] for n in normalised], int(years.pop())


def _convert(df: pd.DataFrame, *, source: str, sheet: str) -> pd.DataFrame:
    """Cast canonical columns to the schema's types, naming any column that fails."""
    types = _schema()["columns"]
    for column in df.columns:
        try:
            values = df[column]
            if not types[column].startswith("string"):
                values = pd.to_numeric(values)
            df[column] = values.astype(types[column])
        except (ValueError, TypeError) as e:
            raise CRDFSchemaError(
                source=source,
                sheet=sheet,
                detail=f"Column {column!r} holds values that are not {types[column]}: {e}",
            ) from e
    return df[[c for c in types if c in df.columns]]


def _read_headers(path: Path) -> tuple[datetime | None, dict[str, list[str]]]:
    """Return a workbook's modified date and the header row of each data sheet."""
    workbook = load_workbook(path, read_only=True)
    try:
        sheets = {}
        for sheet in workbook.worksheets:
            if sheet.title.casefold() == "notes":
                continue
            row = next(sheet.iter_rows(max_row=1, values_only=True), ())
            headers = ["" if v is None else str(v) for v in row]
            while headers and headers[-1] == "":
                headers.pop()
            sheets[sheet.title] = headers
        return workbook.properties.modified, sheets
    finally:
        workbook.close()


def _read_workbook(
    path: Path,
    *,
    perspective: Perspective,
    source_file: str,
    file_etag: str | None,
    pre_process: bool,
) -> pd.DataFrame:
    """Read every data sheet of one CRDF workbook into one DataFrame."""
    modified, sheets = _read_headers(path)

    if modified is None:
        raise CRDFSchemaError(
            source=source_file, detail="The workbook carries no modified date."
        )
    if not sheets:
        raise CRDFSchemaError(
            source=source_file, detail="The workbook has no data sheet."
        )

    types = _schema()["columns"]
    frames = []
    for sheet, headers in sheets.items():
        layout, columns, price_year = _match_layout(
            headers, source=source_file, sheet=sheet
        )
        layout_perspective = _schema()["layouts"][layout]["perspective"]
        if layout_perspective != perspective:
            raise CRDFSchemaError(
                source=source_file,
                sheet=sheet,
                layout=layout,
                detail=f"This is a {layout_perspective} workbook, not {perspective}.",
            )

        # Text columns are read as text so that numeric-looking ids keep their
        # digits.
        df = pd.read_excel(
            path,
            sheet_name=sheet,
            header=None,
            skiprows=1,
            usecols=range(len(headers)),
            dtype={
                i: str for i, c in enumerate(columns) if types[c].startswith("string")
            },
        )
        if pre_process:
            df.columns = columns
            df = _convert(df, source=source_file, sheet=sheet)
        else:
            df.columns = headers
        df["price_year"] = price_year
        frames.append(df)

    df = pd.concat(frames, ignore_index=True)
    df["price_year"] = df["price_year"].astype("int32[pyarrow]")
    df["vintage"] = pd.Series([modified.date()] * len(df), dtype="date32[pyarrow]")
    df["source_file"] = pd.Series(source_file, index=df.index, dtype="string[pyarrow]")
    df["file_etag"] = pd.Series(file_etag, index=df.index, dtype="string[pyarrow]")
    return df


def _files_for_years(
    listing: pd.DataFrame,
    perspective: Perspective,
    start_year: int | None,
    end_year: int | None,
) -> pd.DataFrame:
    """Pick, for each year, the narrowest workbook that holds it.

    Returns:
        The chosen rows of *listing*, with a ``years`` column holding the
        years each workbook is read for. A year appears under one workbook
        only, even when two overlap on it.
    """
    candidates = listing.loc[
        (listing["perspective"] == perspective) & listing["start_year"].notna()
    ]
    available = sorted(
        {
            year
            for start, end in zip(
                candidates["start_year"], candidates["end_year"], strict=True
            )
            for year in range(start, end + 1)
        }
    )
    if not available:
        raise ValueError(f"The CRDF folder at {CRDF_URL} lists no {perspective} files.")

    start = available[0] if start_year is None else start_year
    end = available[-1] if end_year is None else end_year
    if start > end:
        raise ValueError(f"start_year {start} is after end_year {end}.")
    missing = [y for y in range(start, end + 1) if y not in available]
    if missing:
        raise ValueError(
            f"No {perspective} CRDF file holds {missing}. "
            f"Available years: {available[0]} to {available[-1]}."
        )

    ranked = candidates.assign(
        span=candidates["end_year"] - candidates["start_year"]
    ).sort_values(["span", "name"])
    assigned: dict[int, list[int]] = {}
    for year in range(start, end + 1):
        holders = ranked.loc[
            (ranked["start_year"] <= year) & (ranked["end_year"] >= year)
        ]
        assigned.setdefault(holders.index[0], []).append(year)
    return listing.loc[sorted(assigned)].assign(years=pd.Series(assigned))


def download_crdf(
    perspective: Perspective,
    *,
    start_year: int | None = None,
    end_year: int | None = None,
    file: Path | str | None = None,
    pre_process: bool = True,
) -> pd.DataFrame:
    """Download CRDF commitments from the provider or recipient perspective.

    Resolves file names from the live folder listing and reads, for each year,
    the narrowest workbook that holds it (``CRDF-RP_2014-2015.xlsx`` for the
    2015 recipient year). The all-years files are read only when no narrower
    file holds a year.
    Every sheet except ``NOTES`` must match a layout in ``crdf_schema.json``.

    Args:
        perspective: ``"provider"`` or ``"recipient"``.
        start_year: First year to return, inclusive. Defaults to the earliest
            year the folder holds, or every year in *file*.
        end_year: Last year to return, inclusive. Defaults to the latest year
            the folder holds, or every year in *file*.
        file: A local CRDF workbook to read instead of the live folder, such
            as an archived copy. Its layout must still be a known one.
        pre_process: If True (default), rename columns to the schema's names
            and cast them to its types. If False, keep the published headers
            and the types Excel gives.

    Returns:
        pd.DataFrame: One row per CRDF record, amounts in USD thousand, plus
            ``price_year`` (the price year of the constant-price amounts),
            ``vintage`` (the workbook's last-modified date), ``source_file``
            and ``file_etag`` (missing when reading a local *file*).

    Raises:
        CRDFSchemaError: If a sheet's headers match no known layout, the
            layout belongs to the other perspective, or a column fails its type.
        ValueError: If no file holds a requested year.
    """
    if file is not None:
        path = Path(file)
        workbooks = [(path, path.name, None, None)]
    else:
        files = _files_for_years(list_crdf_files(), perspective, start_year, end_year)
        workbooks = []
        for name, url, years in zip(
            files["name"], files["url"], files["years"], strict=True
        ):
            logger.info(f"Reading CRDF workbook {name}")
            path, etag = _fetch(url)
            workbooks.append((path, name, etag, years))

    frames = []
    for path, name, etag, years in workbooks:
        df = _read_workbook(
            path,
            perspective=perspective,
            source_file=name,
            file_etag=etag,
            pre_process=pre_process,
        )
        year = df["year" if pre_process else "Year"]
        # A workbook keeps only the years assigned to it, so a year held by
        # two overlapping workbooks is read once. A local file has no
        # assignment and keeps the requested range.
        if years is not None:
            keep = year.isin(years)
        else:
            keep = pd.Series(True, index=df.index)
            if start_year is not None:
                keep &= year >= start_year
            if end_year is not None:
                keep &= year <= end_year
        frames.append(df.loc[keep])

    return pd.concat(frames, ignore_index=True)
